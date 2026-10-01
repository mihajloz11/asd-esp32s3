"""Gradi TELFOR 2026 rad iz zvanicnog IEEE A4 sablona.

Sablon se otvara, tijelo brise i puni postojecim stilovima (`paper title`,
`Author`, `Abstract`, `Keywords`, `Heading 1..2`, `Body Text`, `table head`,
`figure caption`, `references`). Formatiranje se ne rekonstruise rucno.

Pokretanje iz korijena repoa:
    python radovi/telfor2026/make_figures.py
    python radovi/telfor2026/build_paper.py
"""

from __future__ import annotations

import copy
import sys
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE.parent))
from rezultati import summary  # noqa: E402

TEMPLATE = HERE / "sablon" / "IEEE_conference_template_a4.docx"
OUTPUT = HERE / "telfor2026_asd_esp32s3.docx"

# sadrzaj rada

BANNER_LEFT = "34th Telecommunications Forum TELFOR 2026"
BANNER_RIGHT = "Serbia, Belgrade, November 24-26, 2026"

TITLE = ("Anomalous Sound Detection on an ESP32-S3 with "
         "On-Device Commissioning from Normal Sound")

AUTHORS = "Mihajlo Živković, Ivan Mezei"

AFFIL = [
    "University of Novi Sad, Faculty of Technical Sciences, Novi Sad, Serbia",
    "mihajlo11zivkovic@gmail.com, imezei@uns.ac.rs",
]

ABSTRACT = (
    "This paper presents an anomalous sound detector for fans that runs on an "
    "ESP32-S3 microcontroller with one MEMS microphone. A 96-band log spectrum "
    "from an 8192-point Welch estimate is scored by a squared Mahalanobis "
    "distance, using a Ledoit-Wolf precision matrix learned offline from normal "
    "DCASE 2026 recordings and a center learned on the device. In offline "
    "evaluation on the fan target domain, the design reaches an AUC of 0.863 "
    "with ten calibration windows, "
    "against 0.470 for the evaluated autoencoder baseline, and needs 716 ms of "
    "computation per 10 s window. In the reported trials the device learns the "
    "center, derives entry and release thresholds from 44 normal windows and "
    "checks them on 22 further windows before it starts monitoring. A "
    "reliability gate holds the decision on high-scoring windows whose sub-window spectra "
    "disagree. Physical trials confirmed local commissioning, detection of a "
    "sustained 1 kHz tone and alarm retention under hysteresis. A manually held "
    "paper strip produced variable spectra; one of three blocks raised an "
    "alarm, which was subsequently released. Offline replay reproduced the "
    "reliability and alarm decisions in all 180 monitoring windows. These "
    "results demonstrate the detector's operation and highlight its limitations "
    "under time-varying acoustic conditions and changes in sound level."
)

KEYWORDS = (
    "Anomalous sound detection, condition monitoring, ESP32-S3, Mahalanobis "
    "distance, TinyML, unsupervised learning."
)

# Popunjava se samo zvanicnom oznakom dobijenom iz registracionog sistema.
COPYRIGHT_FOOTER = ""

# fizicke probe, isti izbor prozora kao u master radu
_paper, _tone = summary("paper"), summary("tone")
_pg, _tg = _paper["groups"], _tone["groups"]
_blocks = [_pg[f"airflow_change_paper_{i}"] for i in (1, 2, 3)]


def number(value):
    """Engleski zapis: 1146, 30,255."""
    value = round(value)
    return f"{value:,}" if value >= 10000 else str(value)


def _medians(groups):
    return ", ".join(number(g["median"]) for g in groups)


# ---- tijelo ----
# ('h1'|'h2'|'p'|'table'|'figure', sadrzaj)

BODY = [
    ("h1", "Introduction"),
    ("p",
     "Changes in the sound of a rotating machine can indicate a developing "
     "fault [8], and a "
     "microphone can listen without touching the machine. The DCASE challenge "
     "turned this into an established benchmark: a detector is trained on "
     "normal sound only and has to flag anomalies it has never heard, often on "
     "a machine instance or operating condition that was absent from training "
     "[1]-[5]. Published systems are mostly evaluated offline by the area under "
     "the ROC curve (AUC), on PCs or GPUs rather than on a microcontroller."),
    ("p",
     "A detector mounted next to a machine has a different job. It has to set "
     "a threshold for one particular installation, reject audio it cannot "
     "trust, and decide when a deviation has lasted long enough to report. None"
     " of these steps may rely on fault examples, because a new installation "
     "has none. MLPerf Tiny includes anomaly detection on microcontrollers "
     "and measures inference accuracy, latency and energy [9]. Here the focus "
     "also includes commissioning and alarm decisions for one installation."),
    ("p",
     "This paper describes a complete detector for fans on an ESP32-S3 and what"
     " happened when it was tested on a physical fan. The contributions are: 1)"
     " a statistical detector with higher AUC on the DCASE fan development "
     "data than the "
     "evaluated autoencoder baselines and less than half the measured "
     "computation time of the small int8 autoencoder; 2) commissioning on "
     "the device that learns the center "
     "and both alarm thresholds from normal sound and verifies them before "
     "monitoring; 3) a fail-closed decision chain: a window can raise an "
     "alarm only after it passes quality, machine-presence, reliability and "
     "temporal checks, and a failed check withholds the decision; 4) trials "
     "on a physical fan, with an offline replay of the recorded features, "
     "which confirm that the device made its decisions as specified and show "
     "where detection falls short for unsteady changes."),

    ("h1", "Detector"),
    ("h2", "Hardware and signal path"),
    ("p",
     "The platform is an ESP32-S3-WROOM-1 module (240 MHz, 16 MB PSRAM) running"
     " ESP-IDF v5.5.5. An INMP441 MEMS microphone delivers I2S audio sampled at"
     " 16 kHz and converted to 16-bit samples. The user interface is one "
     "push-button and two LEDs. In the reported trials the button was emulated "
     "over UART through the same event path, and a PC logged telemetry and "
     "operator labels. Fig. 1 shows the signal path. The firmware image, "
     "including research telemetry, is 354,784 B."),
    ("figure", ("FIG_SYSTEM",
                "Signal path and decision chain. Shaded blocks are fail-closed "
                "gates. Calibration, thresholds and alarms are computed on the "
                "device; the PC logs telemetry and sends operator commands.")),
    ("h2", "Spectral feature"),
    ("p",
     "Each 10 s window of 159,744 samples is split into 38 Welch segments [6] of"
     " 8192 samples with a periodic Hann window and 50 % overlap, which gives a"
     " bin spacing of 1.95 Hz. Mean power is taken in 96 logarithmic bands from"
     " 10 Hz to 4 kHz and converted by log10, and the mean over bands is "
     "subtracted to reduce sensitivity to broadband level. The firmware accumulates "
     "segment power as audio arrives, so it never holds the 640 kB float buffer"
     " of a whole window. Fine resolution matters for fans, whose sound is "
     "dominated by narrow lines at the rotation and blade-pass frequencies and "
     "their harmonics [8]. A 1024-point frame has 15.6 Hz spacing, and mel "
     "filtering merges these lines further."),
    ("h2", "Model"),
    ("p",
     "Features are standardized with the mean and standard deviation of 990 "
     "normal source-domain clips of the DCASE 2026 fan [5], and a Ledoit-Wolf "
     "shrinkage estimate [7] of their covariance gives the precision matrix P. "
     "The center c is the mean standardized feature of the calibration windows "
     "recorded on site. A window with standardized feature z receives the score"
     " s = (z − c)ᵀ P (z − c). The covariance describes how the bands of a "
     "normal fan vary together and needs hundreds of clips, while the center "
     "needs only minutes of sound from the actual machine. P occupies 36,864 B "
     "and the normalization 768 B."),
    ("p",
     "The same C source computes features on the PC and on the device. The "
     "maximum absolute feature difference between the Python reference and C "
     "is about 9.54·10⁻⁷ on a benchmark file, and the host-to-device "
     "difference on a live microphone capture is about 1.70·10⁻⁶. These "
     "checks confirm agreement of the feature computation on the tested audio."),
    ("h2", "Offline comparison"),
    ("p",
     "Table I compares the PSD detector design with the DCASE autoencoder "
     "baseline [2] and with log-mel statistics scored by the same Mahalanobis "
     "backend. Local-center results use 100 random calibration splits for "
     "each center size; the two 20-window rows share the same splits. "
     "Calibration windows are never scored. Offline PSD fits use canonical "
     "normalization, while the firmware export adds 10⁻⁸ to the standard "
     "deviation. The autoencoders follow"
     " the DCASE recipe without a local center, so their rows serve only as a "
     "reference point. With the same backend, the PSD "
     "front end gains 0.24 AUC over log-mel statistics. On the device it needs "
     "0.72 s per window, while an int8 autoencoder with 45k parameters on "
     "TensorFlow Lite Micro needs 1.54 s including its log-mel front end."),
    ("table", "T1"),
    ("p",
     "The advantage is specific to the fan. On the other six DCASE machine "
     "types a log-mel front end wins on five, for example 0.539 against 0.448 "
     "on ToyCar, so we claim the result for fans only. Seven further "
     "candidates were declared before evaluation: rotation-order features, "
     "separate centers per operating speed, three variants that use the second"
     " DCASE 2026 channel, transient features, and PSD combined with "
     "transients. None improved on the PSD feature by more than "
     "0.002 AUC over 20 splits. These results come from the development set "
     "that also guided the design, so they should be read as relative "
     "comparisons."),

    ("h1", "Commissioning on the Device"),
    ("h2", "Why ten windows were not enough"),
    ("p",
     "The first firmware set the threshold from the ten calibration windows as "
     "max(Q0.90, mean + 3σ) of their leave-one-out scores. Three calibrations "
     "with the same board, microphone, loudspeaker replay of fan recordings and"
     " room produced thresholds of 5687, 347 and 1088. The thresholds spread by"
     " a factor of 16 while the median score changed by a factor of 4.7, and "
     "92 % of the windows from the first run lay above the threshold of the "
     "second. Ten windows locate the center well, but they say little about "
     "the tail of the normal score distribution that sets the false-alarm "
     "rate."),
    ("h2", "Commissioning sequence"),
    ("p",
     "The firmware estimates the center and the thresholds from separate "
     "windows (Fig. 2). The trials used the GUIDED25 setting: 44 threshold "
     "windows and 22 verification windows, instead of the default 120 and 60. "
     "After at least four settling windows, three of"
     " them stable, the device records ten center windows. The coefficient of "
     "variation of their leave-one-out scores must not exceed 0.6. Up to two "
     "windows with the largest scores may be discarded, recomputing the center "
     "after each; otherwise calibration is rejected. The next 44 windows are "
     "scored against the frozen center. The entry threshold is their empirical"
     " upper 99th percentile, and the release threshold their 95th percentile "
     "kept between the median and half the entry threshold. The last 22 "
     "windows are checked with the same alarm entry and release rule, without "
     "reliability holds. A single alarm episode rejects commissioning. "
     "Including settling, this setting needs at least 80 windows, about "
     "800 s of audio, and took "
     "13.6 min in both trials. The center is never updated afterwards, so "
     "a slowly developing change cannot shift this reference automatically."),
    ("p",
     "The verification step rejected a threshold that was too low. A robust "
     "candidate, the smaller of the empirical Q0.99 and median + 3·1.4826·MAD "
     "with a trimmed center, produced a threshold of 791 in one session. Its "
     "22 verification windows scored from 263 to 7766. Eight exceeded the "
     "threshold, including the last six in a row, so the device correctly "
     "refused to start monitoring."),

    ("h1", "Runtime Decision Chain"),
    ("p",
     "Every window passes four gates before it can raise an alarm. The quality "
     "gate checks the exact sample count, RMS level, peak, clipped fraction (at"
     " most 10⁻³), zero and stuck-sample fractions and a per-window count of "
     "dropped samples, which must be zero, and requires all metrics and "
     "features to be finite. A violation stops the pipeline in every phase "
     "instead of producing a score. The presence gate treats the machine as "
     "running while the window level is no more than 11 dB below the calibrated "
     "level, and as stopped after three consecutive windows below this limit. A "
     "stopped fan can also score highly, so this check prevents low-level "
     "windows from triggering an anomaly alarm."),
    ("p",
     "The reliability gate uses five sub-windows of each 10 s window, formed "
     "from the same Welch segments without extra FFTs. Its instability measure "
     "is the mean over bands of the standard deviation of the five "
     "standardized sub-window features. When a score exceeds the entry "
     "threshold and the instability exceeds 1.25 times the largest value among"
     " the ten calibration windows, the window is put on OBSERVATION_HOLD. A "
     "held window neither counts toward an alarm nor clears an active one, and"
     " it does not change the profile. Six consecutive holds raise a warning. "
     "The gate measures instability without identifying its cause, such as "
     "speech or a knock on the housing."),
    ("p",
     "An alarm requires three consecutive reliable windows above the entry "
     "threshold and is released at the first window at or below the release "
     "threshold. In an earlier comparison of 11 temporal-rule configurations on 2000 "
     "normal windows over 40 splits, a rule of this form, with release at "
     "0.7 of the entry threshold, produced no false alarm episodes, as did "
     "three consecutive windows without hysteresis. On recordings of a "
     "different fan, hysteresis reduced episodes from 11.2 to 5.4 per hour. At the "
     "twelfth alarm window the device emits a sustained-deviation event, and "
     "held windows pause this count. Events are reported as an unknown change,"
     " never as a diagnosed fault."),

    ("h1", "Trials on a Physical Fan"),
    ("p",
     "Both trials used the same firmware image and one electric fan, with the "
     "microphone 40 cm from the fan in a room. In the first, a paper strip "
     "was held by hand in the airflow in three blocks, followed by speech and a door "
     "disturbance. In the second, a loudspeaker played a 1 kHz tone at "
     "varying levels. "
     "The operator labeled conditions from the PC. Metrics use windows with a "
     f"valid protocol, a confirmed label and no label change inside the "
     f"window: {_paper['eligible']} of {_paper['raw']} and {_tone['eligible']} "
     f"of {_tone['raw']} monitoring windows, held windows included. Table II "
     "summarizes both sessions and Fig. 2 shows every window."),
    ("table", "T2"),
    ("figure", ("FIG_TRIALS",
                "Sessions from center learning through monitoring, one point per 10 s "
                "window: (a) paper strip (P), speech (S) and door (D); (b) "
                "1 kHz tone. Shading follows operator labels, not independently "
                "measured stimulus boundaries; the tone continued beyond its "
                "shaded block. Thresholds come from DERIVE. Unconfirmed windows "
                "are excluded from Table II.")),
    ("h2", "Paper strip"),
    ("p",
     "The strip raised the median scores of its three blocks to "
     f"{_medians(_blocks)}, against {number(_pg['normal_baseline']['median'])} "
     "for the baseline, all well above the entry threshold of 8084. In the "
     "first two blocks, however, three of four windows were held, and three "
     "consecutive reliable exceedances never occurred. Holding the strip "
     "steady by hand was difficult. The recorded spectra varied within "
     "and between windows: "
     f"in block one the scores ranged from {number(_blocks[0]['low'])} to "
     f"{number(_blocks[0]['high'])}. "
     "The five sub-window spectra often differed enough to trigger the "
     "reliability gate. An "
     "alarm entered only in block three, lasted two windows into the "
     "following recovery block and was then released. The predefined paper "
     "protocol criteria were not met: only one of three blocks was detected "
     "and the alarm extended into recovery. Speech and door "
     f"windows reached medians of {number(_pg['ambient_speech']['median'])} and "
     f"{number(_pg['ambient_door']['median'])}; all four speech windows and one"
     " of two door windows were held, and neither raised an alarm."),
    ("h2", "Sustained tone"),
    ("p",
     "The tone level varied early in the trial: the 1 kHz band rose from the "
     "first tone window but briefly returned close to its normal level "
     "later. Scores crossed above and below the entry threshold, and several "
     "windows were held. At about 20.5 min both the 1 kHz band and its third "
     "harmonic at 3 kHz rose strongly. This is consistent with a level "
     "increase noted by the operator"
     " without a time stamp. The alarm then "
     "entered after three reliable windows, and the sustained-deviation event "
     "followed 130 s later. The operator marked the end of the tone at "
     "23.3 min, but the 1 kHz band stayed about 42 standardized units above "
     "the center for another 36 windows, indicating continued tonal sound until "
     "about 29.4 min. After that the scores settled between 12,266 and 19,561,"
     " below the entry threshold but above the release threshold of 10,905. "
     "The alarm therefore remained active as specified by hysteresis, which "
     "requires the score to reach the lower threshold rather than a fixed "
     "waiting time. The recording ended 14 windows later without a release, "
     "so its latency was not measured. The trial confirmed tone detection, "
     "sustained-deviation reporting and the expected alarm retention."),
    ("p",
     "Offline verification recomputed the reliability holds from the five "
     "recorded sub-window features and replayed the temporal C code whose "
     "hash matches the trial provenance. All holds and alarm decisions agreed "
     "with the device: 65/65 paper-strip and 115/115 tone windows, including "
     "transitions and unconfirmed labels. This confirms that the recorded "
     "decisions follow the implemented rules; it does not mean that every "
     "induced change was detected. In both "
     "sessions computation took 716 to 728 ms per window, with no reported "
     "sample losses."),
    ("h2", "What the stimuli changed"),
    ("p",
     "Fig. 3 shows the median deviation of the standardized feature from the "
     "center for each stimulus, computed from the feature vectors the device "
     "sent during the trials. The paper strip raised narrow bands between 1.3 "
     "and 3.9 kHz and one near 436 Hz, speech raised the bands from 76 to "
     "210 Hz, and the tone raised the bands at 1 kHz and 3 kHz. For speech the"
     " eight empty bands also dropped by about 12 units, although they contain "
     "no signal at all."),
    ("figure", ("FIG_BANDS",
                "Median deviation z − c of the standardized band feature from "
                "the calibration center, over the eligible windows of each "
                "stimulus (count in parentheses). Ticks mark the eight bands "
                "that contain no FFT bin.")),
    ("h2", "Level dependence of the band map"),
    ("p",
     "At 1.95 Hz spacing, eight of the 96 bands below 28 Hz contain no FFT "
     "bin. The firmware gives them a fixed log value of −20 before the mean is"
     " removed, so they carry the broadband level into all 96 components and "
     "the feature is not level invariant as intended; the speech curve in "
     "Fig. 3 shows this leak. In an offline test that "
     "scales the evaluated clips while keeping the calibration center, a 6 dB "
     "increase multiplies the median normal score by 28 and lowers the AUC "
     "(ten calibration windows, 100 splits) from 0.863 to 0.505."),
    ("p",
     "First, rescoring the original recorded features reproduced the device "
     "scores within a relative error of 10⁻⁶. We then separately reset each "
     "window's band level to the calibration value and rescored it. The "
     "medians of the paper-strip blocks "
     "and of the tone change by less than 4 %, indicating that their high "
     "scores mainly reflect spectral shape. Speech and door medians drop by "
     "26 % and 34 % and remain"
     " several times above the entry threshold."),
    ("p",
     "A band map that advances every band edge by at least one bin keeps 96 "
     "bands over the same range and has no empty bands. It removes this source "
     "of level dependence: on six normal clips and one broadband signal, "
     "the C feature changes by less than 3·10⁻⁶ for gains from 0.25 to 4. "
     "The AUC rises to 0.878, higher in 98 of 100 splits. "
     "It needs a new precision matrix and a new commissioning, and it has not "
     "yet been tested on the fan."),

    ("h1", "Discussion"),
    ("p",
     "The detector targets sustained, reliable spectral deviations; it does "
     "not require a single constant frequency. The hand-held paper strip "
     "produced variable sub-window spectra, whereas the tone eventually "
     "provided three reliable exceedances and a sustained alarm. The gate "
     "responded consistently to this difference. However, it cannot identify "
     "whether instability comes from the machine or its surroundings. The "
     "paper trial therefore demonstrates both the intended rejection of "
     "unstable observations and the limited detection of unsteady changes."),
    ("p",
     "A second limitation is the threshold rule. With the 44 windows used "
     "in these trials, the upper 99th"
     " percentile equals the maximum, so one high normal window sets the entry"
     " threshold. In the tone session a short rise near the end of DERIVE put "
     "it at 21,810, three times the monitoring baseline of 7085, and the early"
     " tone windows scored close to this level. The rule is "
     "conservative on purpose. Two accepted sessions are too few to tell how "
     "repeatable the thresholds are across installations."),

    ("h1", "Conclusion"),
    ("p",
     "A statistical detector with 37.6 kB of fixed model parameters runs on the ESP32-S3 in "
     "716 ms per 10 s window and achieves higher AUC on the DCASE fan "
     "development data than "
     "the evaluated autoencoder baselines. Physical trials demonstrated "
     "commissioning from normal sound, tone detection, sustained-deviation "
     "reporting and alarm retention consistent with hysteresis. The paper "
     "trial demonstrated alarm activation and release in one of three blocks. "
     "Replay reproduced all 180 monitoring decisions, confirming that the "
     "firmware behaved as specified in the tested installation."),
    ("p",
     "The evidence comes from one fan, one room and two sessions, and both "
     "stimuli were induced changes; no mechanical fault was tested. The "
     "paper protocol did not meet its predefined detection and recovery "
     "criteria. Later standalone tests with the push-button and LEDs "
     "confirmed operation as intended, including alarm release after the "
     "tone; only power consumption remains to be measured, with a dedicated "
     "INA226 measurement board. Further work "
     "should repeat the paper trial with a fixed holder, test the "
     "level-invariant band map on hardware and assess "
     "detection of unsteady changes in independent installations."),
]

REFERENCES = [
    "Y. Koizumi et al., “Description and discussion on DCASE2020 Challenge "
    "Task2: Unsupervised anomalous sound detection for machine condition "
    "monitoring,” in Proc. DCASE Workshop, Nov. 2020, pp. 81-85.",
    "N. Harada, D. Niizumi, Y. Ohishi, D. Takeuchi, and M. Yasuda, "
    "“First-shot anomaly sound detection for machine condition monitoring: A "
    "domain generalization baseline,” in Proc. 31st European Signal "
    "Processing Conf. (EUSIPCO), 2023, pp. 191-195, doi: "
    "10.23919/EUSIPCO58844.2023.10289721.",
    "N. Harada et al., “ToyADMOS2: Another dataset of miniature-machine "
    "operating sounds for anomalous sound detection under domain shift "
    "conditions,” in Proc. DCASE Workshop, Barcelona, Spain, Nov. 2021, "
    "pp. 1-5, doi: 10.5281/zenodo.5770113.",
    "K. Dohi et al., “MIMII DG: Sound dataset for malfunctioning "
    "industrial machine investigation and inspection for domain "
    "generalization task,” in Proc. DCASE Workshop, Nancy, France, "
    "Nov. 2022, pp. 1-5.",
    "T. Nishida et al., “Description and discussion on DCASE 2026 "
    "Challenge Task 2: Noise-aware unsupervised anomalous sound detection "
    "for machine condition monitoring,” arXiv:2606.01578, 2026.",
    "P. D. Welch, “The use of the fast Fourier transform for the "
    "estimation of power spectra: A method based on time averaging over "
    "short, modified periodograms,” IEEE Trans. Audio Electroacoust., "
    "vol. 15, no. 2, pp. 70-73, Jun. 1967, doi: "
    "10.1109/TAU.1967.1161901.",
    "O. Ledoit and M. Wolf, “A well-conditioned estimator for "
    "large-dimensional covariance matrices,” J. Multivariate Anal., "
    "vol. 88, no. 2, pp. 365-411, Feb. 2004, doi: "
    "10.1016/S0047-259X(03)00096-4.",
    "R. B. Randall, Vibration-based Condition Monitoring: Industrial, "
    "Aerospace and Automotive Applications. Chichester, U.K.: Wiley, 2011.",
    "C. Banbury et al., “MLPerf Tiny benchmark,” in Proc. NeurIPS Datasets "
    "and Benchmarks Track, 2021.",
]

# ---- tabele ----

TABLES = {
    "T1": {
        "caption": "Detectors on the DCASE 2026 Fan, Target Domain",
        "widths": [3.9, 1.6, 1.9, 1.1],
        "head": ["Detector", "Center", "AUC", "Device"],
        "rows": [
            ["Autoencoder, 270k param. [2]", "none", "0.470 ± 0.010", "-"],
            ["Autoencoder, 45k param., int8", "none", "0.453 ± 0.005", "1.54 s"],
            ["Log-mel statistics + LW", "20 windows", "0.627 ± 0.035", "-"],
            ["PSD, 96 bands + LW", "20 windows", "0.867 ± 0.027", "-"],
            ["PSD, 96 bands + LW", "10 windows", "0.863 ± 0.028", "0.72 s"],
        ],
        "note": "± is the standard deviation over five training seeds for the "
                "autoencoders and over 100 calibration splits otherwise. LW: "
                "Ledoit-Wolf. Device: computation per 10 s window on the "
                "ESP32-S3.",
        "bold_row": 4,
    },
    "T2": {
        "caption": "Physical Fan Trials, Same Firmware Image",
        "widths": [2.9, 2.9, 2.6],
        "head": ["Measurement", "Paper strip", "1 kHz tone"],
        "rows": [
            ["Calibration LOO CV", "0.84 → 0.43, 2 trimmed", "0.43, none trimmed"],
            ["Entry / release threshold", "8084 / 3707", "21,810 / 10,905"],
            ["Baseline median score", number(_pg["normal_baseline"]["median"]),
             number(_tg["normal_baseline"]["median"])],
            ["Stimulus median score", _medians(_blocks),
             number(_tg["constant_tone_1khz"]["median"])],
            ["Held stimulus windows",
             ", ".join(f"{g['hold']}/{g['n']}" for g in _blocks),
             f"{_tg['constant_tone_1khz']['hold']}/{_tg['constant_tone_1khz']['n']}"],
            ["Alarm episodes (device)", "1, block 3; released", "1, retained by hysteresis"],
            ["Sustained-deviation event", "no", "yes"],
            ["Speech / door", "no alarm, 5/6 held", "not applied"],
            ["Paper protocol criteria", "not met: 1/3; carried alarm", "not applicable"],
            ["Decision replay", "65/65 agree", "115/115 agree"],
            ["Dropped samples", "0", "0"],
        ],
        "note": f"Eligible windows: {_paper['eligible']}/{_paper['raw']} and "
                f"{_tone['eligible']}/{_tone['raw']}. Episodes are counted from "
                "the alarm events emitted by the device.",
    },
}

FIGURE_FILES = {
    "FIG_SYSTEM": "slike/fig1_system.png",
    "FIG_TRIALS": "slike/fig2_trials.png",
    "FIG_BANDS": "slike/fig3_bands.png",
}
FIGURE_WIDTH_CM = 8.3


# pomocne funkcije za OOXML

def unnumber(paragraph):
    """Uklanja automatsku numeraciju stila (rimski broj uz Heading 1)."""
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = p_pr.makeelement(qn("w:numPr"), {})
    num_pr.append(num_pr.makeelement(qn("w:ilvl"), {qn("w:val"): "0"}))
    num_pr.append(num_pr.makeelement(qn("w:numId"), {qn("w:val"): "0"}))
    p_pr.append(num_pr)
    return paragraph


def _borders(el_parent, tag, edges, val, sz, color, space="0"):
    borders = el_parent.makeelement(qn(tag), {})
    for edge in edges:
        el = borders.makeelement(qn(f"w:{edge}"), {})
        el.set(qn("w:val"), val.get(edge, "none"))
        el.set(qn("w:sz"), str(sz if val.get(edge, "none") != "none" else 0))
        el.set(qn("w:space"), space)
        el.set(qn("w:color"), color)
        borders.append(el)
    el_parent.append(borders)


def _tight_cells(table):
    """IEEE tabele su zbijene: bez vertikalnog paddinga."""
    tbl_pr = table._tbl.tblPr
    mar = tbl_pr.makeelement(qn("w:tblCellMar"), {})
    for edge, val in (("top", 0), ("bottom", 0), ("left", 57), ("right", 57)):
        el = mar.makeelement(qn(f"w:{edge}"), {})
        el.set(qn("w:w"), str(val))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tbl_pr.append(mar)


def _fixed_layout(table, widths_cm):
    """Sirine u tblGrid, tblW i celijama; LibreOffice cita grid, Word celije."""
    twips = [int(round(w * 567)) for w in widths_cm]
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_pr.append(tbl_pr.makeelement(qn("w:tblLayout"), {qn("w:type"): "fixed"}))
    for tag in ("w:tblW", "w:tblInd"):
        for old in tbl_pr.findall(qn(tag)):
            tbl_pr.remove(old)
    tbl_pr.append(tbl_pr.makeelement(
        qn("w:tblW"), {qn("w:w"): str(sum(twips)), qn("w:type"): "dxa"}))
    for col, w in zip(tbl.tblGrid.findall(qn("w:gridCol")), twips):
        col.set(qn("w:w"), str(w))
    table.autofit = False
    for row in table.rows:
        for cell, w in zip(row.cells, widths_cm):
            cell.width = Cm(w)


def _keep_rows_together(table):
    """Tabela ostaje na jednoj strani i redovi se ne cijepaju."""
    for r, row in enumerate(table.rows):
        tr_pr = row._tr.get_or_add_trPr()
        tr_pr.append(tr_pr.makeelement(qn("w:cantSplit"), {}))
        if r < len(table.rows) - 1:
            for cell in row.cells:
                for p in cell.paragraphs:
                    p.paragraph_format.keep_with_next = True


def add_table(doc, spec):
    cap = doc.add_paragraph(style="table head")
    cap.text = spec["caption"]
    cap.paragraph_format.keep_with_next = True

    head, rows, widths = spec["head"], spec["rows"], spec["widths"]
    table = doc.add_table(rows=len(rows) + 1, cols=len(head))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    _borders(table._tbl.tblPr, "w:tblBorders",
             ("top", "left", "bottom", "right", "insideH", "insideV"),
             {"top": "single", "bottom": "single", "insideH": "single"},
             6, "000000")
    _tight_cells(table)
    _fixed_layout(table, widths)

    for j, text in enumerate(head):
        cell = table.cell(0, j)
        cell.text = ""
        p = cell.paragraphs[0]
        p.style = doc.styles["table col head"]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(text)
        run.bold = True
        run.font.size = Pt(7.5)

    bold_row = spec.get("bold_row")
    for i, row in enumerate(rows):
        for j, text in enumerate(row):
            cell = table.cell(i + 1, j)
            cell.text = ""
            p = cell.paragraphs[0]
            p.style = doc.styles["table copy"]
            run = p.add_run(text)
            run.font.size = Pt(7.5)
            run.bold = i == bold_row
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if j else WD_ALIGN_PARAGRAPH.LEFT

    _keep_rows_together(table)

    if spec.get("note"):
        note = doc.add_paragraph(style="table footnote")
        unnumber(note)
        note.alignment = WD_ALIGN_PARAGRAPH.LEFT
        note.paragraph_format.space_before = Pt(2)
        note.paragraph_format.space_after = Pt(6)
        run = note.add_run(spec["note"])
        run.font.size = Pt(7)
        run.italic = True
    else:
        doc.add_paragraph(style="Body Text")


def add_figure(doc, key, caption):
    path = HERE / FIGURE_FILES[key]
    if not path.exists():
        raise FileNotFoundError(f"{path} ne postoji; pokreni make_figures.py")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(path), width=Cm(FIGURE_WIDTH_CM))
    cap = doc.add_paragraph(style="figure caption")
    cap.text = caption
    doc.add_paragraph(style="Body Text")


# gradnja dokumenta

def build():
    doc = Document(str(TEMPLATE))

    # Sablon nosi tudja svojstva (naslov, autor, urednik); bez ovoga zavrse u PDF-u.
    now = datetime.now()
    cp = doc.core_properties
    cp.title = TITLE
    cp.author = AUTHORS
    cp.last_modified_by = AUTHORS
    cp.keywords = KEYWORDS
    cp.comments = ""
    cp.subject = ""
    cp.category = ""
    cp.created = now
    cp.modified = now
    cp.revision = 1
    # IEEE tipografija ostaje; manji razmaci drze rad u 4 strane.
    doc.styles["Body Text"].paragraph_format.space_after = Pt(5)
    doc.styles["figure caption"].paragraph_format.space_after = Pt(8)
    doc.styles["references"].paragraph_format.space_after = Pt(1)
    body = doc.element.body

    inline_sect_prs = body.xpath("./w:p/w:pPr/w:sectPr")
    title_sect_pr = copy.deepcopy(inline_sect_prs[0])      # 1 kolona, naslov
    body_sect_pr = copy.deepcopy(inline_sect_prs[-1])      # 2 kolone, tijelo

    for child in list(body):
        body.remove(child)

    # naslovni blok, jedna kolona
    banner = doc.add_paragraph()
    banner.paragraph_format.tab_stops.add_tab_stop(Cm(17.8), WD_TAB_ALIGNMENT.RIGHT)
    banner.paragraph_format.space_after = Pt(10)
    banner.add_run(f"{BANNER_LEFT}\t{BANNER_RIGHT}").font.size = Pt(9)

    doc.add_paragraph(TITLE, style="paper title")
    doc.add_paragraph(AUTHORS, style="Author")
    for line in AFFIL:
        p = doc.add_paragraph(line, style="Author")
        for run in p.runs:
            run.font.size = Pt(9)

    closer = doc.add_paragraph()
    closer.paragraph_format.space_after = Pt(0)
    closer._p.get_or_add_pPr().append(title_sect_pr)

    # tijelo, dvije kolone
    for style, label, text in (("Abstract", "Abstract—", ABSTRACT),
                               ("Keywords", "Keywords—", KEYWORDS)):
        p = doc.add_paragraph(style=style)
        r = p.add_run(label)
        r.bold = True
        r.italic = True
        p.add_run(text)

    for kind, payload in BODY:
        if kind == "h1":
            doc.add_paragraph(payload, style="Heading 1")
        elif kind == "h1ack":
            unnumber(doc.add_paragraph(payload, style="Heading 1"))
        elif kind == "h2":
            doc.add_paragraph(payload, style="Heading 2")
        elif kind == "p":
            doc.add_paragraph(payload, style="Body Text")
        elif kind == "table":
            add_table(doc, TABLES[payload])
        elif kind == "figure":
            add_figure(doc, *payload)

    unnumber(doc.add_paragraph("References", style="Heading 1"))
    for ref in REFERENCES:
        doc.add_paragraph(ref, style="references")

    # zavrsna sekcija balansira posljednju stranu
    body.append(body_sect_pr)
    doc.add_section(WD_SECTION.CONTINUOUS)

    # copyright u podnozju prve strane
    section = doc.sections[1]
    section.different_first_page_header_footer = True
    footer = section.first_page_footer
    fp = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    fp.text = ""
    fp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    fp.add_run(COPYRIGHT_FOOTER).font.size = Pt(8)

    doc.save(str(OUTPUT))
    return OUTPUT


if __name__ == "__main__":
    print(f"saved: {build()}")
