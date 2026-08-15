"""Broji preostale [TODO] oznake u radu — tijelo, tabele i podnozja.

Pokretanje:
    ..\\..\\.venv\\Scripts\\python.exe check_todo.py
"""

from __future__ import annotations

import sys
from pathlib import Path

from docx import Document

DOCX = Path(__file__).resolve().parent / "telfor2026_asd_esp32s3.docx"
MARKER = "[TODO]"


def _paragraph_hits(paragraphs, where, hits):
    for p in paragraphs:
        if MARKER in p.text:
            hits.append((where, p.text.strip()))


def main() -> int:
    sys.stdout.reconfigure(encoding="utf-8")
    doc = Document(str(DOCX))
    hits: list[tuple[str, str]] = []

    _paragraph_hits(doc.paragraphs, "tijelo", hits)
    for i, table in enumerate(doc.tables, start=1):
        for row in table.rows:
            for cell in row.cells:
                _paragraph_hits(cell.paragraphs, f"tabela {i}", hits)
    for i, section in enumerate(doc.sections):
        for name in ("first_page_footer", "footer", "first_page_header",
                     "header"):
            part = getattr(section, name)
            if part.is_linked_to_previous:
                continue
            _paragraph_hits(part.paragraphs, f"sekcija {i} / {name}", hits)

    for where, text in hits:
        print(f"{where}: {text[:110]}")
    print(f"\nukupno: {len(hits)}")
    return len(hits)


if __name__ == "__main__":
    main()
