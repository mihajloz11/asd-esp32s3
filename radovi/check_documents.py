"""Check generated manuscripts and save page text for the layout review."""
from __future__ import annotations

import argparse
import json
import re
from pathlib import Path

from docx import Document
from pypdf import PdfReader

ROOT = Path(__file__).resolve().parents[1]
FILES = [
    ("master-rad/master_rad_asd_esp32s3_cir", 6, 20, 12),
    ("master-rad/master_rad_asd_esp32s3_lat", 6, 20, 12),
    ("telfor2026/telfor2026_asd_esp32s3", 3, 2, 9),
]


def check(relative, figures, tables, references):
    base = ROOT / "radovi" / relative
    document = Document(base.with_suffix(".docx"))
    reader = PdfReader(base.with_suffix(".pdf"))
    pages = [page.extract_text() for page in reader.pages]
    full_text = "\n".join(pages)
    cited_paragraphs = []
    for paragraph in document.paragraphs:
        if paragraph.text.strip().upper() in {"REFERENCES", "LITERATURA", "ЛИТЕРАТУРА"}:
            break
        cited_paragraphs.append(paragraph.text)
    citation_text = "\n".join(cited_paragraphs)
    missing = [n for n in range(1, references + 1) if f"[{n}]" not in citation_text]
    # numeric citation ranges also cover the enclosed references
    for first, last in re.findall(r"\[(\d+)\]\s*[-–]\s*\[(\d+)\]", citation_text):
        missing = [n for n in missing if not int(first) <= n <= int(last)]
    result = {
        "file": str(base.relative_to(ROOT)).replace("\\", "/"),
        "pages": len(pages), "figures": len(document.inline_shapes),
        "tables": len(document.tables),
        "missing_citations": missing,
        "todo_occurrences_in_pdf": full_text.count("[TODO]"),
        "unresolved_word_errors": re.findall(r"Error![^\n]*|Грешка![^\n]*", full_text),
        "a4": all(abs(float(p.mediabox.width)-595.276) < 1
                  and abs(float(p.mediabox.height)-841.89) < 1 for p in reader.pages),
        "page_text_lengths": [len(t.strip()) for t in pages],
    }
    result["checks_pass"] = (result["a4"] and not missing
                              and not result["unresolved_word_errors"]
                              and result["figures"] == figures
                              and result["tables"] >= tables
                              and min(result["page_text_lengths"]) > 10
                              and (len(pages) <= 4 if "telfor" in relative else 30 <= len(pages) <= 50))
    preview = base.parent / "preview"
    preview.mkdir(exist_ok=True)
    (preview / (base.name + "_pages.txt")).write_text(
        "\n\n".join(f"PAGE {i}\n{text}" for i, text in enumerate(pages, 1)), encoding="utf-8")
    return result


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("only", nargs="*", help="dio putanje, npr. telfor2026")
    args = parser.parse_args()
    result = [check(*entry) for entry in FILES
              if not args.only or any(part in entry[0] for part in args.only)]
    args.output.write_text(json.dumps(result, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    for item in result:
        print(json.dumps({k:v for k,v in item.items() if k != "page_text_lengths"}, ensure_ascii=False))
    return 0 if all(r["checks_pass"] for r in result) else 1


if __name__ == "__main__":
    raise SystemExit(main())
